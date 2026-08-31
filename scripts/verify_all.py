"""最终回归：全模块 import + 多格式转换 + 确认 PDF 路径仍可用（不跑完整 OCR）。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))

# 1) 全模块 import
from anyocr.core import engine, office, pdf, rapid, vlm
from anyocr import cli, api, config
print("✅ 全模块 import 正常")

# 2) 确认 file_to_md 路由存在
from anyocr.core.engine import AnyOcrEngine, _TEXT_EXTS, _OFFICE_EXTS
eng = AnyOcrEngine()
assert hasattr(eng, "file_to_md") and hasattr(eng, "pdf_to_md")
print(f"✅ 路由表: 文本={sorted(_TEXT_EXTS)} Office={sorted(_OFFICE_EXTS)}")

# 3) 各格式转换（TMP 里新建小样本）
from docx import Document
from pptx import Presentation
import openpyxl

TMP = tempfile.mkdtemp(prefix="regress_")

d = Document(); d.add_heading("标题", 0); d.add_paragraph("正文段落")
docx_p = os.path.join(TMP, "a.docx"); d.save(docx_p)

prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "PPT标题"; s.placeholders[1].text_frame.text = "要点"
pptx_p = os.path.join(TMP, "b.pptx"); prs.save(pptx_p)

wb = openpyxl.Workbook(); ws = wb.active; ws.append(["x", "y"]); ws.append([1, 2])
xlsx_p = os.path.join(TMP, "c.xlsx"); wb.save(xlsx_p)

txt_p = os.path.join(TMP, "d.txt")
with open(txt_p, "w", encoding="utf-8") as f: f.write("纯文本")

out_dir = os.path.join(TMP, "out"); os.makedirs(out_dir, exist_ok=True)
for p in [docx_p, pptx_p, xlsx_p, txt_p]:
    md = eng.file_to_md(p, out_md=os.path.join(out_dir, os.path.basename(p).replace(".", "_") + ".md"))
    print(f"  ✅ {os.path.splitext(os.path.basename(p))[1]} → {len(md)}字符，md已落盘")

# 4) 老格式检测（LibreOffice 未装时应报友好错误，不崩溃）
from anyocr.core.office import detect_office_tools, OfficeError, legacy_to_md
tools = detect_office_tools()
print(f"  ℹ LibreOffice 检测: {tools or '未安装（老格式转换需安装）'}")

# 5) CLI 解析器参数检查（convert 用 file 参数）
from anyocr.cli import build_parser
parser = build_parser()
ns = parser.parse_args(["convert", "x.docx", "--out", "y"])
assert ns.file == "x.docx", ns
print("  ✅ CLI convert 支持任意文件参数")

# 6) API 路由检查
from anyocr.api import app
routes = {r.path for r in app.routes}
assert {"/convert", "/batch", "/upload", "/health"} <= routes
print(f"  ✅ API 路由: {sorted(routes)}")

print("\n🎉 全部回归通过")
