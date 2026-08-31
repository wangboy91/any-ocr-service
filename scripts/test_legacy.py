"""老格式端到端测试：生成 docx → soffice 转 .doc → legacy_to_md 解析回 md。"""
import os
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "src"))

from anyocr.core.office import detect_office_tools, office_to_md

print("1) 检测 LibreOffice 工具:", detect_office_tools() or "❌ 未找到")

tools = detect_office_tools()
if not tools:
    sys.exit("LibreOffice 不可用，无法测试老格式")
soffice = tools.get("soffice") or tools.get("libreoffice")
print(f"   使用: {soffice}")

TMP = tempfile.mkdtemp(prefix="legacy_test_")

# 2) 生成 docx 样例
from docx import Document
d = Document()
d.add_heading("老格式转换测试", 0)
d.add_heading("正文标题", level=1)
d.add_paragraph("这是一段测试中文，验证 doc 老格式转换链路。")
d.add_paragraph("第二个要点，测试多段落。")
docx_p = os.path.join(TMP, "source.docx")
d.save(docx_p)
print(f"\n2) 已生成 docx: {docx_p}")

# 3) soffice headless 转成 .doc（老二进制格式）
doc_p = os.path.join(TMP, "source.doc")
cmd = [soffice, "--headless", "--convert-to", "doc", "--outdir", TMP, docx_p]
print(f"\n3) soffice 转换 docx→doc: {' '.join(cmd)}")
proc = subprocess.run(cmd, capture_output=True, text=True, timeout=240)
print(f"   返回码: {proc.returncode}, 存在: {os.path.exists(doc_p)}")
if not os.path.exists(doc_p):
    print("   soffice 输出:", (proc.stdout or proc.stderr)[-300:])
    sys.exit("doc 转换失败")

# 4) 用 legacy_to_md 解析 .doc → md
print(f"\n4) 解析 .doc → Markdown:")
md = office_to_md(doc_p)
print(f"   输出 {len(md)} 字符")
for line in md.splitlines()[:12]:
    print("   |", line)

print("\n🎉 老格式 doc 转换链路测试完成")
