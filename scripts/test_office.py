"""测试多格式转换：生成 docx/pptx/xlsx/txt 样例 → file_to_md → 打印结果片段。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))

from anyocr.core.engine import AnyOcrEngine

TMP = tempfile.mkdtemp(prefix="office_test_")
print("测试目录:", TMP)

# ---- 生成样例文件 ----
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
import openpyxl

# DOCX
d = Document()
d.add_heading("测试文档标题", 0)
d.add_heading("第一章 背景", level=1)
d.add_paragraph("这是一段正文，用于验证 docx 解析是否正常。")
d.add_paragraph("带编号列表项一", style="List Bullet")
t = d.add_table(rows=2, cols=2)
t.rows[0].cells[0].text = "列A"
t.rows[0].cells[1].text = "列B"
t.rows[1].cells[0].text = "值1"
t.rows[1].cells[1].text = "值2"
docx_path = os.path.join(TMP, "示例文档.docx")
d.save(docx_path)

# PPTX
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "演示文稿标题"
body = slide.placeholders[1].text_frame
body.text = "第一点内容"
body.add_paragraph().text = "第二点内容"
tbl_slide = prs.slides.add_slide(prs.slide_layouts[5])
pptx_path = os.path.join(TMP, "示例演示.pptx")
prs.save(pptx_path)

# XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "数据表"
ws.append(["姓名", "分数"])
ws.append(["张三", 95])
ws.append(["李四", 87])
xlsx_path = os.path.join(TMP, "示例表格.xlsx")
wb.save(xlsx_path)

# TXT
txt_path = os.path.join(TMP, "示例笔记.txt")
with open(txt_path, "w", encoding="utf-8") as f:
    f.write("这是一段纯文本笔记。\n第二行内容。")

eng = AnyOcrEngine()
for p in [docx_path, pptx_path, xlsx_path, txt_path]:
    print("\n" + "=" * 50)
    print(f"▶ {os.path.basename(p)}")
    try:
        md = eng.file_to_md(p)
        print(f"  ✅ 输出 {len(md)} 字符")
        for line in md.splitlines()[:14]:
            print("   |", line)
    except Exception as e:
        print(f"  ❌ 失败: {e}")

print("\n" + "=" * 50)
print("全部测试完成")
